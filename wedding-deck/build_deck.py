# -*- coding: utf-8 -*-
"""結婚式二次会用 ○×クイズ＆ビンゴ進行スライドを組み立てる。

構成とテンション感は、お預かりした『チーム対抗企画スライド 3』を下敷きにしている：
  問題 → クリックで10秒カウントダウン（BGM付き）→ クリックで正解発表（効果音＋赤枠点滅）
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

import content
from pptx_motion import AudioLibrary, Timeline, apply_timing, apply_transition

ROOT = Path(__file__).parent
IMG = ROOT / "assets" / "img"
SFX = ROOT / "assets" / "sfx"
OUT = ROOT / "build" / "結婚式_マルバツクイズ＆ビンゴ.pptx"

SLIDE_W, SLIDE_H = 12192000, 6858000

# ---- 配色 ----------------------------------------------------------------
INK = RGBColor(0x3A, 0x2E, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xFD, 0xF7, 0xF2)
ROSE = RGBColor(0xD9, 0x4F, 0x70)        # ○ 側
NAVY = RGBColor(0x34, 0x55, 0x8B)        # × 側
GOLD = RGBColor(0xC9, 0xA2, 0x27)
DEEP = RGBColor(0x96, 0x2C, 0x42)
ALERT = RGBColor(0xFF, 0x00, 0x00)

JP = "Meiryo"

# ---- レイアウト（EMU） ---------------------------------------------------
BADGE = dict(x=620000, y=300000, cx=2500000, cy=820000)
TIMER = dict(x=10400000, y=250000, cx=1300000, cy=1300000)
Q_TEXT = dict(x=900000, y=1620000, cx=10392000, cy=1500000)
PANEL = dict(y=3300000, cy=2050000, cx=4250000, lx=1150000, rx=6792000)
RIBBON = dict(x=1150000, y=5550000, cx=9892000, cy=900000)
MARK = 1000000          # ○×マークの一辺

# 写真ありレイアウト
PH_PHOTO = dict(x=700000, y=1250000, cx=3000000, cy=4497000)
PH_TEXT = dict(x=4100000, y=520000, cx=5900000, cy=1300000)
PH_PANEL = dict(cx=7300000, x=4100000, y1=2000000, y2=3650000, cy=1450000)
PH_RIBBON = dict(x=4100000, y=5350000, cx=7300000, cy=1100000)
PH_MARK = 900000


def _next_id(slide):
    return slide.shapes._spTree.max_shape_id + 1


def add_bg(slide, name):
    pic = slide.shapes.add_picture(str(IMG / name), 0, 0, SLIDE_W, SLIDE_H)
    slide.shapes._spTree.remove(pic._element)
    slide.shapes._spTree.insert(2, pic._element)   # 背景として最背面へ
    return pic


def textbox(slide, x, y, cx, cy, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    """lines は (文字列, ポイント, 太字, 色) のリスト。改行は \n で表現する。"""
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(cx), Emu(cy))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(90000)
    tf.margin_top = tf.margin_bottom = Emu(45000)
    first = True
    for text, size, bold, color in lines:
        for i, chunk in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            p.line_spacing = 1.18
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.size = Pt(size)
            f.bold = bold
            f.color.rgb = color
            f.name = JP
            r.font._rPr.set("lang", "ja-JP")
            r.font._rPr.set("altLang", "en-US")
    return box


def rounded(slide, x, y, cx, cy, fill, line, line_w=4.5, radius=0.12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(cx), Emu(cy))
    sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line
    sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def fill_shape_text(shape, lines, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_top = tf.margin_bottom = Emu(45000)
    first = True
    for text, size, bold, color in lines:
        for chunk in text.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            p.line_spacing = 1.05
            r = p.add_run()
            r.text = chunk
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = JP


def add_mark(slide, kind, cx_center, cy_center, size):
    """○と×を文字ではなく図形で描く。会場の後ろの席からでも読めるように。"""
    def paint(sh, color):
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    if kind == "maru":
        x, y = cx_center - size // 2, cy_center - size // 2
        ring = slide.shapes.add_shape(MSO_SHAPE.DONUT, Emu(x), Emu(y), Emu(size), Emu(size))
        ring.adjustments[0] = 0.13
        return paint(ring, ROSE)

    # × は「プラス図形を回転」だと太りすぎるので、棒 2 本を交差させて描く
    length, thick = int(size * 0.94), int(size * 0.20)
    last = None
    for rot in (45, 135):
        bar = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(cx_center - length // 2), Emu(cy_center - thick // 2), Emu(length), Emu(thick),
        )
        bar.adjustments[0] = 0.5
        bar.rotation = rot
        last = paint(bar, NAVY)
    return last


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# --------------------------------------------------------------------------
def add_timer(slide):
    """10秒カウントダウンの丸と、0〜10 の数字テキストを重ねて置く。"""
    oval = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Emu(TIMER["x"]), Emu(TIMER["y"]), Emu(TIMER["cx"]), Emu(TIMER["cy"])
    )
    oval.fill.solid()
    oval.fill.fore_color.rgb = DEEP
    oval.line.color.rgb = WHITE
    oval.line.width = Pt(5)
    oval.shadow.inherit = False

    numbers = []
    for n in range(10, -1, -1):
        box = slide.shapes.add_textbox(
            Emu(TIMER["x"]), Emu(TIMER["y"] + 190000), Emu(TIMER["cx"]), Emu(900000)
        )
        tf = box.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(n)
        r.font.size = Pt(54)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = "Arial"
        numbers.append(box.shape_id)
    return oval.shape_id, numbers


def quiz_slide(prs, audio, index, q):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    photo = q.get("layout") == "photo"
    add_bg(slide, "bg_quiz.png")

    # 第N問バッジ
    badge = rounded(slide, BADGE["x"], BADGE["y"], BADGE["cx"], BADGE["cy"], DEEP, GOLD, 3, 0.5)
    label = "最終問題" if q.get("final") else f"第 {index} 問"
    fill_shape_text(badge, [(label, 26, True, WHITE)])

    if photo:
        slide.shapes.add_picture(
            str(IMG / "groom_child.jpg"),
            Emu(PH_PHOTO["x"]), Emu(PH_PHOTO["y"]),
            Emu(PH_PHOTO["cx"]), Emu(PH_PHOTO["cy"]),
        )
        textbox(slide, PH_TEXT["x"], PH_TEXT["y"], PH_TEXT["cx"], PH_TEXT["cy"],
                [(q["statement"], 27, True, INK)])
        geo = dict(
            lx=PH_PANEL["x"], ly=PH_PANEL["y1"], rx=PH_PANEL["x"], ry=PH_PANEL["y2"],
            cx=PH_PANEL["cx"], cy=PH_PANEL["cy"],
        )
        ribbon_geo, mark_size, stacked = PH_RIBBON, PH_MARK, True
    else:
        textbox(slide, Q_TEXT["x"], Q_TEXT["y"], Q_TEXT["cx"], Q_TEXT["cy"],
                [(q["statement"], 33, True, INK)])
        geo = dict(
            lx=PANEL["lx"], ly=PANEL["y"], rx=PANEL["rx"], ry=PANEL["y"],
            cx=PANEL["cx"], cy=PANEL["cy"],
        )
        ribbon_geo, mark_size, stacked = RIBBON, MARK, False

    maru = rounded(slide, geo["lx"], geo["ly"], geo["cx"], geo["cy"], WHITE, ROSE, 6)
    batsu = rounded(slide, geo["rx"], geo["ry"], geo["cx"], geo["cy"], WHITE, NAVY, 6)
    for kind, panel, side in (("maru", maru, "左"), ("batsu", batsu, "右")):
        if stacked:
            add_mark(slide, kind, panel.left + 780000, panel.top + panel.height // 2, mark_size)
            textbox(slide, panel.left + 1450000, panel.top, panel.width - 1750000, panel.height,
                    [(f"画面に向かって {side} へ移動", 24, True, INK)], align=PP_ALIGN.LEFT)
        else:
            add_mark(slide, kind, panel.left + panel.width // 2,
                     panel.top + 640000, mark_size)
            textbox(slide, panel.left, panel.top + 1290000, panel.width, 640000,
                    [(f"画面に向かって {side} へ移動", 19, True, INK)])

    oval_id, number_ids = add_timer(slide)

    # 正解の赤枠（クリック2で出現して点滅する）
    correct = maru if q["answer"] == "maru" else batsu
    pad = 110000
    frame = rounded(
        slide,
        correct.left - pad, correct.top - pad,
        correct.width + pad * 2, correct.height + pad * 2,
        None, ALERT, 7.5,
    )

    ribbon = rounded(slide, ribbon_geo["x"], ribbon_geo["y"], ribbon_geo["cx"],
                     ribbon_geo["cy"], DEEP, GOLD, 3, 0.25)
    fill_shape_text(ribbon, [(q["reveal"], 20 if photo else 23, True, WHITE)])

    sid = _next_id(slide)
    cd_id, cd_dur = audio.add(slide, str(SFX / "sfx_countdown10.mp3"), "カウントダウンBGM", sid, 0)
    ans_id, ans_dur = audio.add(slide, str(SFX / "sfx_answer.mp3"), "正解発表効果音", sid + 1, 1)

    tl = Timeline()
    tl.group().countdown(oval_id, number_ids, cd_id, cd_dur)
    tl.group().play(ans_id, ans_dur).appear(frame.shape_id).pulse(frame.shape_id) \
              .appear(ribbon.shape_id)
    apply_transition(slide, "fade")
    apply_timing(slide, tl)

    ans_jp = "○" if q["answer"] == "maru" else "×"
    add_notes(slide, "\n".join([
        f"【第{index}問】正解：{ans_jp}",
        "① 問題を読み上げる",
        "② クリック → 10秒カウントダウン（BGMが鳴ります）。会場の皆さまに移動していただく",
        "③ クリック → 正解の赤枠が点滅＋効果音、解説が出ます",
        "",
        q.get("note", ""),
    ]).strip())
    return slide


def section_slide(prs, audio, head, sub=None, bg="bg_section.png", sfx=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    textbox(slide, 900000, 2150000, 10392000, 1450000, [(head, 56, True, WHITE)])
    if sub:
        textbox(slide, 900000, 3750000, 10392000, 700000,
                [(sub, 22, False, RGBColor(0xF6, 0xDF, 0xC8))])
    if sfx:
        sid = _next_id(slide)
        spid, dur = audio.add(slide, str(SFX / sfx), "セクション効果音", sid)
        tl = Timeline()
        tl.group(auto=True).play(spid, dur)   # スライドを開いた瞬間に鳴る
        apply_transition(slide, "fade")
        apply_timing(slide, tl)
    else:
        apply_transition(slide, "fade")
    return slide


def bullet_slide(prs, head, title, lines, bg="bg_quiz.png", move=False, big=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, bg)
    badge = rounded(slide, BADGE["x"], BADGE["y"], BADGE["cx"], BADGE["cy"], DEEP, GOLD, 3, 0.5)
    fill_shape_text(badge, [(head, 24, True, WHITE)])
    textbox(slide, 900000, 1250000, 10392000, 1000000, [(title, 40, True, DEEP)])

    if move:
        # ○＝左／×＝右 を大きな図で示す
        for kind, px, arrow in (("maru", 1150000, "← 画面に向かって 左 へ"),
                                ("batsu", 6792000, "画面に向かって 右 へ →")):
            panel = rounded(slide, px, 2600000, 4250000, 2350000, WHITE,
                            ROSE if kind == "maru" else NAVY, 6)
            add_mark(slide, kind, px + 2125000, 3300000, 1150000)
            textbox(slide, px, 4180000, 4250000, 640000, [(arrow, 21, True, INK)])
        textbox(slide, 900000, 5150000, 10392000, 700000,
                [("※ 迷ったときは、まわりの流れに乗ってしまってOKです🙆", 18, True, INK)])
    elif big:
        textbox(slide, 900000, 2400000, 10392000, 2600000, [(big, 60, True, DEEP)])
    else:
        body = [(f"・{t}", 25, False, INK) for t in lines]
        textbox(slide, 1300000, 2350000, 9600000, 3400000, body,
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
    return slide


def bingo_slide(prs, audio, item):
    slide = bullet_slide(
        prs, "ビンゴ大会", item["title"], item.get("lines", []),
        bg="bg_bingo.png", big=item.get("big"),
    )
    sid = _next_id(slide)
    name = "ドラムロール" if item["sfx"] == "drumroll" else "ファンファーレ"
    spid, dur = audio.add(slide, str(SFX / f"sfx_{item['sfx']}.mp3"), name, sid)
    tl = Timeline()
    tl.group(auto=True).play(spid, dur)   # スライドが切り替わるたびに音が鳴る
    apply_transition(slide, "push")
    apply_timing(slide, tl)
    add_notes(slide, f"スライドを表示すると自動で「{name}」が鳴ります。")
    return slide


def bgm_slide(prs, audio):
    """『家族になろうよ』をクリックで再生するスライド。"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, "bg_title.png")
    textbox(slide, 900000, 900000, 10392000, 1100000, [("BGM", 28, True, GOLD)])
    textbox(slide, 900000, 1850000, 10392000, 1300000,
            [("家族になろうよ", 54, True, DEEP)])

    part, dur = audio.add(
        slide, str(SFX / "bgm_kazoku_ni_narouyo_PLACEHOLDER.mp3"),
        "家族になろうよ", _next_id(slide),
    )
    # 音声アイコンは画面外なので、押しやすい再生ボタンを別に置いてトリガーにする
    btn = rounded(slide, 4096000, 3350000, 4000000, 1300000, DEEP, GOLD, 4, 0.3)
    btn._element._nvXxPr.cNvPr.set("name", "BGM再生ボタン")
    fill_shape_text(btn, [("▶　BGMを再生", 32, True, WHITE)])
    textbox(slide, 900000, 4950000, 10392000, 1200000,
            [("↑ このボタン自体をクリックすると再生されます（スライド送りでは鳴りません）", 18, True, INK),
             ("（今は無音のダミー音源が入っています。mp3をいただければ差し替えます）", 16, False, INK)])

    tl = Timeline()
    tl.trigger(btn.shape_id).play(part, dur)   # ▶ボタンを押したときだけ再生
    apply_transition(slide, "fade")
    apply_timing(slide, tl)
    add_notes(slide, "\n".join([
        "【差し替え手順（ご自身でやる場合）】",
        "1. このスライドの左外（グレー部分）にあるスピーカーアイコンを選択して削除",
        "2. [挿入] → [オーディオ] → [このコンピューター上のオーディオ] で「家族になろうよ」のmp3を挿入",
        "3. 挿入したスピーカーアイコンをスライドの外へドラッグして見えない位置に移動",
        "4. 挿入した音声を選択 →[アニメーション]→[アニメーションウィンドウ]→[トリガー]→",
        "   [クリック時] →「BGM再生ボタン」 を指定する",
        "   （※もともと同じ設定が入っています。同じ形にすれば動きは変わりません）",
        "",
        "※ mp3をお送りいただければ、この作業込みで組み込んだものをお渡しできます。",
    ]))
    return slide


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    audio = AudioLibrary(prs, str(IMG / "audio_icon.png"))

    # --- 表紙 ---
    title = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(title, "bg_title.png")
    textbox(title, 900000, 1500000, 10392000, 900000,
            [("新郎新婦のことどれだけ知ってる？", 26, True, GOLD)])
    textbox(title, 900000, 2400000, 10392000, 1600000,
            [("○× クイズ大会", 68, True, DEEP)])
    textbox(title, 900000, 4200000, 10392000, 900000,
            [("＆ ビンゴ大会 🎉", 30, True, INK)])
    apply_transition(title, "fade")

    # --- ルール説明 ---
    for r in content.RULES:
        bullet_slide(prs, r["head"], r["title"], r.get("lines", []), move=r.get("move", False))

    # --- ○×クイズ ---
    section_slide(prs, audio, "○×クイズ", "さあ、スタートです！", sfx="sfx_answer.mp3")
    for i, q in enumerate(content.QUIZ, start=1):
        quiz_slide(prs, audio, i, q)

    # --- 結果発表 ---
    section_slide(prs, audio, "結果発表", "最後まで勝ち残ったのは…？", sfx="sfx_drumroll.mp3")
    win = bullet_slide(prs, "結果発表", "優勝！", [], big="{{優勝者のお名前}} 様")
    apply_transition(win, "fade")
    add_notes(win, "{{優勝者のお名前}} を当日その場で入力するか、口頭で読み上げてください。")

    # --- ビンゴ大会 ---
    section_slide(prs, audio, "ビンゴ大会", "お手元のカードをご用意ください🎱",
                  bg="bg_section.png", sfx="sfx_drumroll.mp3")
    for item in content.BINGO:
        bingo_slide(prs, audio, item)

    # --- エンディング ---
    bgm_slide(prs, audio)
    end = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(end, "bg_title.png")
    textbox(end, 900000, 1900000, 10392000, 1600000,
            [("本日はご列席いただき\nありがとうございました", 42, True, DEEP)])
    textbox(end, 900000, 3900000, 10392000, 1200000,
            [("これからもふたりをよろしくお願いします🤍", 24, True, INK)])
    apply_transition(end, "fade")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"{OUT}  ({OUT.stat().st_size/1024/1024:.1f} MB, {len(prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
